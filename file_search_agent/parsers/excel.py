"""
Excel Parser - Extracts data from Excel files.
Uses openpyxl for .xlsx and xlrd for legacy .xls files.
"""

from pathlib import Path
from typing import List, Dict, Any, Optional
import pandas as pd
from openpyxl import load_workbook

from .base import BaseParser, ParsedDocument, DocumentChunk


class ExcelParser(BaseParser):
    """Parser for Excel spreadsheets."""
    
    supported_extensions = [".xlsx", ".xls", ".xlsm"]
    
    def can_parse(self, file_path: Path) -> bool:
        return file_path.suffix.lower() in self.supported_extensions
    
    def parse(self, file_path: Path) -> ParsedDocument:
        """Parse Excel file and convert sheets to searchable text."""
        file_path = Path(file_path)
        
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        # Read all sheets
        excel_data = pd.read_excel(file_path, sheet_name=None, dtype=str)
        
        # Get workbook metadata (xlsx only)
        metadata = {}
        if file_path.suffix.lower() == ".xlsx":
            try:
                wb = load_workbook(file_path, read_only=True)
                props = wb.properties
                metadata = {
                    "creator": props.creator,
                    "title": props.title,
                    "subject": props.subject,
                    "description": props.description,
                }
                wb.close()
            except Exception:
                pass
        
        # Convert sheets to text
        all_content = []
        chunks = []
        chunk_index = 0
        char_position = 0
        
        for sheet_name, df in excel_data.items():
            sheet_text = self._dataframe_to_text(df, sheet_name)
            all_content.append(f"=== Sheet: {sheet_name} ===\n{sheet_text}")
            
            # Create chunks for this sheet
            sheet_chunks = self.chunk_text(
                sheet_text,
                metadata={"sheet_name": sheet_name}
            )
            
            for chunk in sheet_chunks:
                chunk.chunk_index = chunk_index
                chunk.start_char += char_position
                chunk.end_char += char_position
                chunk.metadata["sheet_name"] = sheet_name
                chunks.append(chunk)
                chunk_index += 1
            
            char_position += len(sheet_text) + len(f"=== Sheet: {sheet_name} ===\n") + 2
        
        full_content = "\n\n".join(all_content)
        file_meta = self.get_file_metadata(file_path)
        
        return ParsedDocument(
            file_path=file_path,
            file_name=file_path.name,
            file_type="excel",
            content=full_content,
            chunks=chunks,
            file_size=file_meta["file_size"],
            created_at=file_meta.get("created_at"),
            modified_at=file_meta.get("modified_at"),
            title=metadata.get("title") or file_path.stem,
            author=metadata.get("creator"),
            page_count=len(excel_data),  # Number of sheets
            metadata={
                "sheet_names": list(excel_data.keys()),
                "total_rows": sum(len(df) for df in excel_data.values()),
                **metadata
            }
        )
    
    def _dataframe_to_text(self, df: pd.DataFrame, sheet_name: str) -> str:
        """Convert a DataFrame to searchable text format."""
        if df.empty:
            return f"[Empty sheet: {sheet_name}]"
        
        # Fill NaN values
        df = df.fillna("")
        
        lines = []
        
        # Add column headers
        headers = " | ".join(str(col) for col in df.columns)
        lines.append(f"Headers: {headers}")
        lines.append("-" * 40)
        
        # Add each row as readable text
        for idx, row in df.iterrows():
            # Format: "Column1: Value1, Column2: Value2, ..."
            row_parts = []
            for col in df.columns:
                value = str(row[col]).strip()
                if value:  # Only include non-empty values
                    row_parts.append(f"{col}: {value}")
            
            if row_parts:
                lines.append("; ".join(row_parts))
        
        return "\n".join(lines)
    
    def parse_as_dataframes(self, file_path: Path) -> Dict[str, pd.DataFrame]:
        """Return raw DataFrames for programmatic access."""
        return pd.read_excel(file_path, sheet_name=None)
    
    def search_in_excel(
        self,
        file_path: Path,
        query: str,
        case_sensitive: bool = False
    ) -> List[Dict[str, Any]]:
        """Search for a value across all sheets and return matching cells."""
        results = []
        dfs = self.parse_as_dataframes(file_path)
        
        for sheet_name, df in dfs.items():
            df_str = df.astype(str)
            
            if not case_sensitive:
                query = query.lower()
                df_str = df_str.apply(lambda x: x.str.lower())
            
            # Find matching cells
            mask = df_str.apply(lambda col: col.str.contains(query, na=False))
            
            for row_idx, row in mask.iterrows():
                for col_name in row.index:
                    if row[col_name]:
                        results.append({
                            "sheet": sheet_name,
                            "row": row_idx + 2,  # +2 for 1-based + header
                            "column": col_name,
                            "value": str(df.loc[row_idx, col_name])
                        })
        
        return results
